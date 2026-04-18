#!/usr/bin/env python3
"""Generate a PowerPoint rationalization deck for the INSIGHT project."""

import json
from pathlib import Path
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData

# ── Colors ──
DARK_BG = RGBColor(0x0F, 0x0F, 0x1A)
CARD_BG = RGBColor(0x16, 0x21, 0x3E)
BORDER = RGBColor(0x1E, 0x2D, 0x50)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x9C, 0xA3, 0xAF)
LIGHT_GRAY = RGBColor(0xD1, 0xD5, 0xDB)
RED = RGBColor(0xEF, 0x44, 0x44)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
YELLOW = RGBColor(0xEA, 0xB3, 0x08)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
CYAN = RGBColor(0x06, 0xB6, 0xD4)
PURPLE = RGBColor(0xA8, 0x55, 0xF7)


def set_slide_bg(slide, color=DARK_BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=14,
                 color=WHITE, bold=False, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Calibri"
    p.alignment = alignment
    return txBox


def add_stat_card(slide, left, top, label, value, color=BLUE):
    width, height = 2.1, 1.0
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_top = Pt(8)
    tf.margin_left = Pt(10)
    p = tf.paragraphs[0]
    p.text = str(value)
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = color
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.size = Pt(9)
    p2.font.color.rgb = GRAY


def add_table(slide, left, top, width, col_widths, headers, rows):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_shape = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top),
                                          Inches(width), Inches(0.3 * n_rows))
    table = table_shape.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(9)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.font.name = "Calibri"
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x1A, 0x2A, 0x4A)
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(8)
                p.font.color.rgb = LIGHT_GRAY
                p.font.name = "Calibri"
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 0 else RGBColor(0x12, 0x1A, 0x2E)
    return table_shape


def main():
    output_dir = Path("INSIGHT/rationalization")

    # Load data
    analysis = json.loads((output_dir / "analysis" / "summary.json").read_text(encoding="utf-8"))
    classification = json.loads((output_dir / "analysis" / "classification.json").read_text(encoding="utf-8"))
    manifest = json.loads((output_dir / "run_manifest.json").read_text(encoding="utf-8"))
    active = json.loads((output_dir / "telemetry" / "active_reports.json").read_text(encoding="utf-8"))
    retire = json.loads((output_dir / "telemetry" / "retire_reports.json").read_text(encoding="utf-8"))
    high_impact = json.loads((output_dir / "analysis" / "high_impact.json").read_text(encoding="utf-8"))
    duplicates = json.loads((output_dir / "analysis" / "duplicate_sql.json").read_text(encoding="utf-8"))
    orphans = json.loads((output_dir / "analysis" / "orphans.json").read_text(encoding="utf-8"))
    unused_metrics = json.loads((output_dir / "analysis" / "unused_metrics.json").read_text(encoding="utf-8"))
    stale = json.loads((output_dir / "analysis" / "stale_objects.json").read_text(encoding="utf-8"))

    dedup_info = {}
    dedup_path = output_dir / "inventory" / "active_reports_dedup.json"
    if dedup_path.exists():
        dedup_info = json.loads(dedup_path.read_text(encoding="utf-8"))

    sql_reports = []
    sql_path = output_dir / "inventory" / "active_reports_sql.json"
    if sql_path.exists():
        sql_reports = json.loads(sql_path.read_text(encoding="utf-8"))

    # Full inventory
    full_inv = {}
    for cat in ("reports", "documents", "cubes", "metrics", "filters", "prompts",
                "attributes", "facts", "tables", "security_filters"):
        f = output_dir / "inventory" / f"{cat}.json"
        if f.exists():
            full_inv[cat] = len(json.loads(f.read_text(encoding="utf-8")))
    total_inv = sum(full_inv.values())

    total_reports = full_inv.get("reports", classification["totalInventoryReports"])
    project_id = manifest.get("project", {}).get("id", "")

    has_sql = sum(1 for r in sql_reports if r.get("sql"))
    dedup_families = dedup_info.get("familyCount", 0)
    dedup_skipped = dedup_info.get("duplicateCount", 0)
    unique_active = classification["uniqueActiveReports"]

    sql_report_ids = {r["id"] for r in sql_reports}
    active_dupe_groups = []
    for g in duplicates:
        members = [o for o in g["objects"] if o["id"] in sql_report_ids]
        if len(members) >= 2:
            active_dupe_groups.append(members)
    sql_reducible = sum(len(g) for g in active_dupe_groups) - len(active_dupe_groups)
    after_sql_dedup = unique_active - sql_reducible

    # AST analysis
    ast_path = output_dir / "analysis" / "ast_duplicates.json"
    ast_groups = 0
    ast_removable = 0
    ast_parse_failures = 0
    if ast_path.exists():
        ast_data = json.loads(ast_path.read_text(encoding="utf-8"))
        ast_groups = ast_data.get("exactAstGroups", 0)
        ast_removable = ast_data.get("removableReports", 0)
        ast_parse_failures = ast_data.get("parseFailures", 0)

    structural_removable = max(sql_reducible, ast_removable)
    after_structural_dedup = unique_active - structural_removable

    # Semantic review
    reviews_path = output_dir / "sql" / "cluster_reviews.json"
    semantic_removable = 0
    semantic_reviews = []
    semantic_actions = {}
    if reviews_path.exists():
        rev_data = json.loads(reviews_path.read_text(encoding="utf-8"))
        semantic_reviews = rev_data.get("reviews", [])
        for rev in semantic_reviews:
            a = rev.get("analysis", {})
            if "error" in a:
                continue
            semantic_removable += a.get("removable_count", 0)
            act = a.get("action", "?")
            semantic_actions[act] = semantic_actions.get(act, 0) + 1

    final_unique = after_structural_dedup - semantic_removable

    orphan_cats = {}
    for o in orphans:
        orphan_cats[o.get("category", "?")] = orphan_cats.get(o.get("category", "?"), 0) + 1

    # Top active
    top_active = sorted(active, key=lambda x: -(x.get("totalExecutions") or 0))[:10]

    # ════════════════════════════════════════════════════════════════
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ── SLIDE 1: Title ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 1, 1.5, 11, 1.2,
                 "MicroStrategy Report Rationalization", 36, WHITE, True, PP_ALIGN.CENTER)
    add_text_box(slide, 1, 2.8, 11, 0.8,
                 "INSIGHT Project", 28, BLUE, False, PP_ALIGN.CENTER)
    add_text_box(slide, 1, 3.8, 11, 0.5,
                 f"Analysis Date: {datetime.now().strftime('%B %d, %Y')}", 14, GRAY, False, PP_ALIGN.CENTER)
    add_text_box(slide, 1, 4.5, 11, 0.5,
                 f"{total_reports:,} reports analyzed  |  {total_inv:,} total objects  |  "
                 f"{analysis['totalEdges']:,} dependency edges",
                 14, GRAY, False, PP_ALIGN.CENTER)
    add_text_box(slide, 1, 6.2, 11, 0.5,
                 "Ralph Lauren - Analytics Sandbox Environment", 12,
                 RGBColor(0x6B, 0x72, 0x80), False, PP_ALIGN.CENTER)

    # ── SLIDE 2: Executive Summary ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "Executive Summary", 28, WHITE, True)

    add_stat_card(slide, 0.5, 1.2, "Total Reports", f"{total_reports:,}", BLUE)
    add_stat_card(slide, 2.8, 1.2, "Active (Retain)", f"{len(active):,}", GREEN)
    add_stat_card(slide, 5.1, 1.2, "Retire", f"{len(retire):,}", RED)
    add_stat_card(slide, 7.4, 1.2, "Final Unique", f"{final_unique:,}", CYAN)
    add_stat_card(slide, 9.7, 1.2, "Reduction",
                  f"{(total_reports - final_unique)/total_reports*100:.0f}%", ORANGE)

    add_text_box(slide, 0.5, 2.5, 12, 0.4, "Key Findings:", 16, WHITE, True)
    findings = [
        f"{len(retire):,} reports ({len(retire)/total_reports*100:.0f}%) have no execution evidence and should be retired",
        f"{dedup_families:,} families of definition-identical reports found ({dedup_skipped:,} copies removed)",
        f"{len(active_dupe_groups):,} groups of reports with identical SQL ({sql_reducible:,} reducible)",
        f"Semantic (LLM) analysis identified {len(semantic_reviews):,} near-duplicate clusters ({semantic_removable:,} further reducible)",
        f"{analysis['unusedMetricCount']:,} of {classification['totalMetrics']:,} metrics ({analysis['unusedMetricCount']/classification['totalMetrics']*100:.0f}%) are unused",
        f"Final unique report count: {final_unique:,} (from {total_reports:,} — {(total_reports - final_unique)/total_reports*100:.1f}% reduction)",
    ]
    y = 3.0
    for f in findings:
        add_text_box(slide, 0.8, y, 11, 0.35, f"  \u2022  {f}", 12, LIGHT_GRAY)
        y += 0.38

    # ── SLIDE 3: Reduction Funnel ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "Reduction Funnel", 28, WHITE, True)
    add_text_box(slide, 0.5, 0.9, 12, 0.4,
                 f"How we went from {total_reports:,} reports to {final_unique:,}", 14, GRAY)

    funnel = [
        ["Original Inventory", f"{total_reports:,}", "", "All reports enumerated via REST API"],
        ["After Telemetry Retirement", f"{len(active):,}", f"-{len(retire):,}",
         "No execution evidence in 7 months"],
        ["After Definition Dedup", f"{unique_active:,}", f"-{dedup_skipped:,}",
         f"{dedup_families:,} families of identical definitions"],
        ["After Structural Dedup (SQL+AST)", f"{after_structural_dedup:,}", f"-{structural_removable:,}",
         f"{len(active_dupe_groups):,} hash groups + {ast_groups:,} AST groups"],
        ["After Semantic (LLM) Dedup", f"{final_unique:,}", f"-{semantic_removable:,}",
         f"{len(semantic_reviews):,} near-duplicate clusters reviewed"],
    ]
    add_table(slide, 0.5, 1.6, 12, [4, 2, 1.5, 4.5],
              ["Stage", "Reports", "Change", "Explanation"], funnel)

    # Funnel chart
    chart_data = CategoryChartData()
    chart_data.categories = ["Inventory", "After Telemetry", "After Def Dedup",
                              "After Structural Dedup", "After Semantic Dedup"]
    chart_data.add_series("Reports",
                          [total_reports, len(active), unique_active, after_structural_dedup, final_unique])
    chart_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(4.2), Inches(12), Inches(2.8),
        chart_data
    )
    chart = chart_frame.chart
    chart.has_legend = False
    series = chart.plots[0].series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = BLUE
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.color.rgb = LIGHT_GRAY
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.tick_labels.font.color.rgb = GRAY

    # ── SLIDE 4: Inventory Breakdown ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "Full Inventory Breakdown", 28, WHITE, True)

    cats_sorted = sorted(full_inv.items(), key=lambda x: -x[1])
    inv_rows = [[cat.replace("_", " ").title(), f"{count:,}", f"{count/total_inv*100:.1f}%"]
                for cat, count in cats_sorted if count > 0]
    inv_rows.append(["TOTAL", f"{total_inv:,}", "100%"])
    add_table(slide, 0.5, 1.2, 5, [2.5, 1.5, 1],
              ["Category", "Count", "%"], inv_rows)

    chart_data = CategoryChartData()
    chart_data.categories = [c for c, _ in cats_sorted if _ > 0]
    chart_data.add_series("Objects", [n for _, n in cats_sorted if n > 0])
    chart_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, Inches(6), Inches(1.2), Inches(6.5), Inches(5),
        chart_data
    )
    chart = chart_frame.chart
    chart.has_legend = False
    chart.plots[0].series[0].format.fill.solid()
    chart.plots[0].series[0].format.fill.fore_color.rgb = BLUE
    chart.category_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.color.rgb = LIGHT_GRAY
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.tick_labels.font.color.rgb = GRAY

    # ── SLIDE 5: Telemetry Results ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "Telemetry Matching Results", 28, WHITE, True)

    add_stat_card(slide, 0.5, 1.2, "Telemetry Records", "87,274", BLUE)
    add_stat_card(slide, 2.8, 1.2, "Active (Matched)", f"{len(active):,}", GREEN)
    add_stat_card(slide, 5.1, 1.2, "Retire (Unmatched)", f"{len(retire):,}", RED)
    add_stat_card(slide, 7.4, 1.2, "Match Rate",
                  f"{len(active)/total_reports*100:.1f}%", YELLOW)

    # Pie
    pie_data = CategoryChartData()
    pie_data.categories = ["Active (Retain)", "Retire"]
    pie_data.add_series("Reports", [len(active), len(retire)])
    pie_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(0.5), Inches(2.8), Inches(5.5), Inches(4),
        pie_data
    )
    pie = pie_frame.chart
    pie.has_legend = True
    pie.legend.position = XL_LEGEND_POSITION.BOTTOM
    pie.legend.font.size = Pt(11)
    pie.legend.font.color.rgb = LIGHT_GRAY
    pie.plots[0].series[0].points[0].format.fill.solid()
    pie.plots[0].series[0].points[0].format.fill.fore_color.rgb = GREEN
    pie.plots[0].series[0].points[1].format.fill.solid()
    pie.plots[0].series[0].points[1].format.fill.fore_color.rgb = RED

    # Top active table
    add_text_box(slide, 6.5, 2.5, 6, 0.4, "Top Active Reports by Execution:", 14, WHITE, True)
    top_rows = [[r.get("name", "")[:40], f"{r.get('totalExecutions', 0):,}",
                 str(r.get("totalUsers", 0))]
                for r in top_active]
    add_table(slide, 6.5, 3.0, 6, [3.5, 1.5, 1],
              ["Report Name", "Executions", "Users"], top_rows)

    # ── SLIDE 6: Definition Dedup ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "Definition-Based De-duplication", 28, WHITE, True)

    add_stat_card(slide, 0.5, 1.2, "Active Reports", f"{len(active):,}", BLUE)
    add_stat_card(slide, 2.8, 1.2, "Families Found", f"{dedup_families:,}", PURPLE)
    add_stat_card(slide, 5.1, 1.2, "Copies Removed", f"{dedup_skipped:,}", RED)
    add_stat_card(slide, 7.4, 1.2, "Unique After Dedup", f"{unique_active:,}", GREEN)

    add_text_box(slide, 0.5, 2.5, 12, 0.8,
                 "Method: Generated a fingerprint for each report by hashing its attributes, metrics, "
                 "filter text, source type, and source cube ID. Reports with identical fingerprints are "
                 "true copies regardless of name. The copy with the highest execution count was kept.",
                 12, GRAY)

    if dedup_info.get("families"):
        add_text_box(slide, 0.5, 3.5, 12, 0.4, "Largest Families:", 14, WHITE, True)
        fam_rows = []
        for fam in dedup_info["families"][:12]:
            fam_rows.append([
                fam.get("parentName", "")[:45],
                str(fam.get("familySize", 0)),
                str(fam.get("familySize", 0) - 1),
            ])
        add_table(slide, 0.5, 4.0, 8, [4.5, 1.5, 2],
                  ["Parent Report", "Family Size", "Copies Removed"], fam_rows)

    # ── SLIDE 7: SQL & Duplicates ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "SQL Extraction & Duplicate Analysis", 28, WHITE, True)

    add_stat_card(slide, 0.5, 1.2, "SQL Extracted", f"{has_sql:,}", GREEN)
    add_stat_card(slide, 2.8, 1.2, "SQL Failed", f"{len(sql_reports) - has_sql:,}", RED)
    add_stat_card(slide, 5.1, 1.2, "Dup SQL Groups", f"{len(active_dupe_groups):,}", YELLOW)
    add_stat_card(slide, 7.4, 1.2, "SQL Reducible", f"{sql_reducible:,}", ORANGE)

    if active_dupe_groups:
        add_text_box(slide, 0.5, 2.5, 12, 0.4, "Largest Duplicate SQL Groups:", 14, WHITE, True)
        dup_rows = []
        for i, members in enumerate(sorted(active_dupe_groups, key=len, reverse=True)[:12], 1):
            names = [m.get("name", "")[:35] for m in members]
            dup_rows.append([
                str(i), str(len(members)),
                names[0],
                ", ".join(n[:25] for n in names[1:3]) + ("..." if len(names) > 3 else ""),
            ])
        add_table(slide, 0.5, 3.0, 12, [0.8, 1, 4, 6.2],
                  ["#", "Size", "Primary Report", "Also includes"], dup_rows)

    # ── SLIDE 7a: AST Structural Analysis ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "AST Structural Analysis (sqlglot)", 28, WHITE, True)
    add_text_box(slide, 0.5, 0.9, 12, 0.5,
                 "SQL parsed into Abstract Syntax Trees via sqlglot; Merkle subtree hashes computed. "
                 "Catches structural duplicates that hash-based dedup misses (alias order, JOIN order).",
                 12, GRAY)

    add_stat_card(slide, 0.5, 1.7, "Reports with SQL", f"{has_sql:,}", BLUE)
    add_stat_card(slide, 2.8, 1.7, "Parsed OK", f"{has_sql - ast_parse_failures:,}", GREEN)
    add_stat_card(slide, 5.1, 1.7, "AST Dup Groups", f"{ast_groups:,}", YELLOW)
    add_stat_card(slide, 7.4, 1.7, "Removable", f"{ast_removable:,}", ORANGE)

    add_text_box(slide, 0.5, 3.0, 12, 0.4,
                 f"Comparison: SQL hash dedup found {sql_reducible:,} removable; "
                 f"AST dedup found {ast_removable:,} removable "
                 f"({'+' + str(ast_removable - sql_reducible) if ast_removable > sql_reducible else 'same or fewer'} structural duplicates).",
                 13, LIGHT_GRAY)

    add_text_box(slide, 0.5, 4.0, 12, 0.4, "How it works:", 14, WHITE, True)
    steps_text = [
        "1. Parse each SQL into an Abstract Syntax Tree using sqlglot (Redshift dialect)",
        "2. Canonicalize: drop literals, normalize column references, sort commutative nodes",
        "3. Compute Merkle-style hash for every subtree in the AST",
        "4. Group reports by their full subtree-hash set — identical sets = structural duplicates",
    ]
    y = 4.5
    for s in steps_text:
        add_text_box(slide, 0.8, y, 11, 0.3, s, 11, LIGHT_GRAY)
        y += 0.35

    add_text_box(slide, 0.5, 6.5, 12, 0.5,
                 f"Result: {ast_groups:,} exact-AST groups. Using max(SQL hash, AST) = "
                 f"{structural_removable:,} structural duplicates going forward.",
                 11, GRAY)

    # ── SLIDE 7b: Semantic Analysis (LLM-powered) ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "Semantic Analysis (LLM-Powered)", 28, WHITE, True)
    add_text_box(slide, 0.5, 0.9, 12, 0.5,
                 "Hash-based detection catches only identical SQL. Embeddings + LLM review find "
                 "near-duplicates with different surface syntax.",
                 12, GRAY)

    add_stat_card(slide, 0.5, 1.7, "Clusters Reviewed", f"{len(semantic_reviews):,}", PURPLE)
    add_stat_card(slide, 2.8, 1.7, "PARAMETERIZE", f"{semantic_actions.get('PARAMETERIZE', 0):,}", ORANGE)
    add_stat_card(slide, 5.1, 1.7, "MERGE_IMMEDIATE", f"{semantic_actions.get('MERGE_IMMEDIATE', 0):,}", RED)
    add_stat_card(slide, 7.4, 1.7, "Removable Reports", f"{semantic_removable:,}", GREEN)

    add_text_box(slide, 0.5, 3.0, 12, 0.4, "Top Consolidation Opportunities:", 14, WHITE, True)
    by_removable = sorted(
        [r for r in semantic_reviews if "error" not in r.get("analysis", {})],
        key=lambda r: -r.get("analysis", {}).get("removable_count", 0)
    )[:12]
    sem_rows = []
    for rev in by_removable:
        a = rev["analysis"]
        sem_rows.append([
            str(rev["clusterId"]),
            a.get("label", "?")[:35],
            a.get("action", "?"),
            str(rev["totalReportCount"]),
            str(a.get("removable_count", 0)),
        ])
    if sem_rows:
        add_table(slide, 0.5, 3.5, 12, [1.5, 4.5, 2.5, 1.5, 2],
                  ["Cluster", "Description", "Action", "Reports", "Remove"], sem_rows)

    add_text_box(slide, 0.5, 6.7, 12, 0.5,
                 "Method: all-MiniLM-L6-v2 embeddings + agglomerative clustering (85% similarity) + GPT-4o review. "
                 "Selected based on discrimination test (separation score 0.288 vs OpenAI 0.162, Voyage 0.047).",
                 10, GRAY)

    # ── SLIDE 8: Orphans & Unused Metrics ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "Orphans & Unused Metrics", 28, WHITE, True)

    add_stat_card(slide, 0.5, 1.2, "Orphan Objects", f"{len(orphans):,}", YELLOW)
    add_stat_card(slide, 2.8, 1.2, "Unused Metrics",
                  f"{analysis['unusedMetricCount']:,}", ORANGE)
    add_stat_card(slide, 5.1, 1.2, "Total Metrics", f"{classification['totalMetrics']:,}", BLUE)
    add_stat_card(slide, 7.4, 1.2, "% Unused",
                  f"{analysis['unusedMetricCount']/max(classification['totalMetrics'],1)*100:.0f}%", RED)

    # Orphan breakdown
    add_text_box(slide, 0.5, 2.5, 6, 0.4, "Orphans by Category:", 14, WHITE, True)
    orph_sorted = sorted(orphan_cats.items(), key=lambda x: -x[1])
    orph_rows = [[cat, f"{count:,}"] for cat, count in orph_sorted]
    orph_rows.append(["TOTAL", f"{len(orphans):,}"])
    add_table(slide, 0.5, 3.0, 5, [2.5, 2.5],
              ["Category", "Orphan Count"], orph_rows)

    # Orphan chart
    if orph_sorted:
        orph_data = CategoryChartData()
        orph_data.categories = [c for c, _ in orph_sorted]
        orph_data.add_series("Orphans", [n for _, n in orph_sorted])
        orph_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED, Inches(6.5), Inches(2.5), Inches(6), Inches(4),
            orph_data
        )
        orph_chart = orph_frame.chart
        orph_chart.has_legend = False
        orph_chart.plots[0].series[0].format.fill.solid()
        orph_chart.plots[0].series[0].format.fill.fore_color.rgb = YELLOW
        orph_chart.category_axis.tick_labels.font.color.rgb = LIGHT_GRAY
        orph_chart.value_axis.tick_labels.font.color.rgb = GRAY

    # ── SLIDE 9: High-Impact & Dependencies ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "Dependency Graph & High-Impact Objects", 28, WHITE, True)

    add_stat_card(slide, 0.5, 1.2, "Dependency Edges", f"{analysis['totalEdges']:,}", CYAN)
    add_stat_card(slide, 2.8, 1.2, "High-Impact (Top 50)", f"{analysis['highImpactCount']}", RED)
    add_stat_card(slide, 5.1, 1.2, "Stale Objects", f"{analysis['staleObjectCount']:,}", YELLOW)
    add_stat_card(slide, 7.4, 1.2, "Stale & Active", f"{classification['staleAndActive']:,}", ORANGE)

    add_text_box(slide, 0.5, 2.5, 12, 0.4, "Top 15 High-Impact Objects:", 14, WHITE, True)
    hi_rows = []
    for h in high_impact[:15]:
        hi_rows.append([
            h.get("name", "")[:45],
            h.get("category", ""),
            str(h.get("dependentCount", 0)),
        ])
    add_table(slide, 0.5, 3.0, 10, [5, 2.5, 2.5],
              ["Object Name", "Category", "Dependents"], hi_rows)

    # ── SLIDE 10: Recommendations ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "Recommendations", 28, WHITE, True)

    recs = [
        (f"1. Retire {len(retire):,} unused reports",
         f"No execution evidence in 7 months. Move to archive folder. "
         f"This alone is an {len(retire)/total_reports*100:.0f}% reduction.", RED),
        (f"2. Remove {dedup_skipped:,} definition-identical copies",
         f"{dedup_families:,} families of identical reports. Keep the highest-execution copy, "
         f"archive the rest.", RED),
        (f"3. Consolidate {sql_reducible:,} SQL-duplicate reports",
         f"{len(active_dupe_groups):,} groups of reports producing identical queries. "
         f"Replace with parameterized reports.", ORANGE),
        (f"4. Consolidate {semantic_removable:,} semantic near-duplicates (LLM-identified)",
         f"{semantic_actions.get('PARAMETERIZE', 0)} PARAMETERIZE + "
         f"{semantic_actions.get('MERGE_IMMEDIATE', 0)} MERGE_IMMEDIATE clusters. "
         f"Replace with parameterized master reports.", ORANGE),
        (f"5. Clean up {analysis['unusedMetricCount']:,} unused metrics",
         f"{analysis['unusedMetricCount']/classification['totalMetrics']*100:.0f}% of all metrics "
         f"are not used by any active report. Simplify the project schema.", YELLOW),
        ("6. Protect high-impact objects",
         "Top objects have hundreds of dependents. Changes must be carefully validated.", GREEN),
        ("7. Schedule quarterly reviews",
         "Prevent future report sprawl with recurring rationalization.", BLUE),
    ]
    y = 1.2
    for title, body, color in recs:
        add_text_box(slide, 0.5, y, 12, 0.3, title, 14, color, True)
        add_text_box(slide, 0.8, y + 0.35, 11, 0.45, body, 11, LIGHT_GRAY)
        y += 0.85

    # ── SLIDE 11: Next Steps ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_text_box(slide, 0.5, 0.3, 12, 0.6, "Next Steps", 28, WHITE, True)

    steps = [
        "1.  Share retirement list with business stakeholders for 2-week review",
        f"2.  Begin bulk retirement of {len(retire):,} unused reports (lowest risk)",
        f"3.  Remove {dedup_skipped:,} definition-identical copies",
        f"4.  Consolidate {len(active_dupe_groups):,} SQL-duplicate groups",
        f"5.  Review {len(semantic_reviews):,} LLM-identified near-duplicate clusters (target: {semantic_removable:,} reports)",
        f"6.  Clean up {analysis['unusedMetricCount']:,} unused metrics and {analysis['orphanCount']:,} orphans",
        "7.  Schedule quarterly rationalization reviews",
    ]
    y = 1.3
    for step in steps:
        add_text_box(slide, 0.8, y, 11, 0.4, step, 14, LIGHT_GRAY)
        y += 0.55

    add_text_box(slide, 0.5, 5.8, 12, 0.5,
                 f"Total reduction: {total_reports:,} -> {final_unique:,} reports "
                 f"({(total_reports - final_unique)/total_reports*100:.1f}% reduction)",
                 18, GREEN, True, PP_ALIGN.CENTER)

    # Save
    out_path = output_dir / "INSIGHT_Rationalization_Deck.pptx"
    prs.save(str(out_path))
    print(f"Deck saved to: {out_path}")


if __name__ == "__main__":
    main()
